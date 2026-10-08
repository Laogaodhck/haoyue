"""DOCX 渲染器：基于 python-docx，按块类型映射标题/列表/段落样式。"""

from __future__ import annotations

from pathlib import Path

from ..core.base import BaseRenderer
from ..core.errors import DocumentError
from ..core.models import BlockType, Document, DocumentBlock


class DocxRenderer(BaseRenderer):
    name = "docx"
    extensions = frozenset({"docx"})
    default_suffix = ".docx"

    def write(self, document: Document, out_path: Path) -> Path:
        try:
            from docx import Document as DocxDocument
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError as exc:  # pragma: no cover
            raise DocumentError(
                "DOCX 输出依赖 python-docx 未安装，请执行：pip install python-docx",
                detail=str(exc),
            ) from exc

        docx = DocxDocument()
        for block in document.blocks:
            para = docx.add_paragraph()
            match block.type:
                case BlockType.HEADING:
                    # docx 标题级别 1-9，钳制到有效范围。
                    para.style = docx.styles[f"Heading {max(1, min(9, block.level))}"]
                    para.add_run(block.text)
                case BlockType.LIST_ITEM:
                    style_name = "List Number" if block.ordered else "List Bullet"
                    para.style = docx.styles[style_name]
                    # 嵌套层级用前导缩进模拟，避免依赖复杂的编号定义。
                    if block.indent:
                        para.paragraph_format.left_indent = None
                        para.paragraph_format.first_line_indent = None
                        para.text = "  " * block.indent + block.text
                    else:
                        para.add_run(block.text)
                case BlockType.QUOTE:
                    para.style = docx.styles["Intense Quote"]
                    para.add_run(block.text)
                case BlockType.CODE:
                    run = para.add_run(block.text)
                    run.font.name = "Consolas"
                case BlockType.DIVIDER:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.add_run("────────")
                case _:  # 未知类型降级为段落，保证转换不中断
                    _append_multiline(para, block.text)
        docx.save(str(out_path))
        return out_path


def _append_multiline(para, text: str) -> None:
    """多行文本在 DOCX 中用软换行（w:br）而非硬拆多段。"""
    lines = text.split("\n")
    para.add_run(lines[0])
    for line in lines[1:]:
        para.add_run().add_break()
        para.add_run(line)
