"""渲染器保真度验证：六种块类型 × docx/xlsx/pptx 三种目标格式。

每个用例先把含全部六种块类型的 Markdown 渲染为目标格式，再从产物
读回断言（python-docx / openpyxl / python-pptx），保证「写进去的
内容真的以正确语义落盘」而非仅文件存在。

用法::

    python doc_toolkit/verify_renderers.py
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from doc_toolkit import default_toolkit
from doc_toolkit.core.errors import DocumentError

SAMPLE_MD = """# 季度评审报告

## 背景说明

本季度完成了运行时重构与本地推理加速两项重点工作。

- 重构事件溯源层，崩溃可对账
- CUDA 后端解码提速 11 倍
  - KV 前缀复用进一步压低首字延迟

1. 提交评审材料
2. 归档评测报告

> 所有改动均通过 696 项单测与 70% 覆盖率门槛。

```python
def add(a, b):
    return a + b
```

---

| 指标 | 基线 | 优化后 |
|---|---|---|
| 解码 tok/s | 7.6 | 84.9 |
| TTFT | 20.5s | 0.61s |
"""

CHECKS: list[tuple[str, bool]] = []


def check(name: str, ok: bool) -> None:
    CHECKS.append((name, ok))
    print(f"  {'✔' if ok else '✘'} {name}")


def verify_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    check("docx: 标题落 Heading 1 样式", "Heading 1" in styles and "季度评审报告" in texts)
    check("docx: 二级标题落 Heading 2 样式", "Heading 2" in styles and "背景说明" in texts)
    check("docx: 段落文本保留", any("运行时重构" in t for t in texts))
    check("docx: 无序列表落 List Bullet 样式",
          "List Bullet" in styles and any("事件溯源层" in t for t in texts))
    check("docx: 有序列表落 List Number 样式", "List Number" in styles)
    check("docx: 引用块落 Intense Quote 样式",
          "Intense Quote" in styles and any("696 项单测" in t for t in texts))
    check("docx: 代码块使用等宽字体",
          any(run.font.name == "Consolas" for p in doc.paragraphs for run in p.runs))
    check("docx: 分隔线渲染", any("────" in t for t in texts))


def verify_xlsx(path: Path) -> None:
    from openpyxl import load_workbook

    sheet = load_workbook(str(path)).active
    rows = [[("" if c.value is None else str(c.value)) for c in row] for row in sheet.iter_rows()]
    flat = [cell for row in rows for cell in row]
    check("xlsx: 标题行落分区标题", "季度评审报告" in flat)
    check("xlsx: 段落文本落单列行", any("运行时重构" in cell for cell in flat))
    check("xlsx: Markdown 表格切分为真实单元格",
          any(row[:3] == ["指标", "基线", "优化后"] for row in rows))
    check("xlsx: 表头分隔行被跳过", not any(cell.startswith("-") and cell.endswith("-") for row in rows for cell in row))
    check("xlsx: 表格数据行切分正确", any(row[:3] == ["解码 tok/s", "7.6", "84.9"] for row in rows))
    check("xlsx: 引用块带注释前缀", any(cell.startswith("# ") for cell in flat))
    check("xlsx: 代码块逐行落格", "def add(a, b):" in flat)
    bold_rows = [row for row in sheet.iter_rows()
                 if any(cell.font.bold for cell in row)]
    check("xlsx: 分区标题加粗", any(any(
        cell.value == "季度评审报告" and cell.font.bold for cell in row) for row in bold_rows))


def verify_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides = list(prs.slides)
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in slides]
    check("pptx: 一级标题生成独立幻灯片", "季度评审报告" in titles)
    check("pptx: 幻灯片数量随标题分页", len(slides) >= 1)
    all_text = "\n".join(
        shape.text_frame.text
        for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    check("pptx: 段落进入正文", "运行时重构" in all_text)
    check("pptx: 列表项进入正文", "事件溯源层" in all_text)
    check("pptx: 引用块进入正文", "696 项单测" in all_text)
    check("pptx: 代码块进入正文", "def add(a, b):" in all_text)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="doc-toolkit-renderers-") as tmp:
        source = Path(tmp) / "sample.md"
        source.write_text(SAMPLE_MD, encoding="utf-8")

        cases = [
            ("docx", verify_docx),
            ("xlsx", verify_xlsx),
            ("pptx", verify_pptx),
        ]
        for fmt, verifier in cases:
            print(f"[{fmt}]")
            try:
                out = default_toolkit.convert(source, fmt)
            except DocumentError as exc:
                check(f"{fmt}: 渲染（{exc}）", False)
                continue
            check(f"{fmt}: 产物生成 {out.name}", out.is_file() and out.stat().st_size > 0)
            verifier(out)

    failed = [name for name, ok in CHECKS if not ok]
    print(f"\n通过 {len(CHECKS) - len(failed)}/{len(CHECKS)} 项。")
    if failed:
        print("失败项：")
        for name in failed:
            print(f"  - {name}")
        return 1
    print("渲染器保真度验证全部通过。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
