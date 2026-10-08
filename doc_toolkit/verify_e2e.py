"""doc_toolkit 端到端验证脚本（一次性，不作为包的一部分发布）。"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
from doc_toolkit import (  # noqa: E402
    CorruptedDocumentError,
    DocumentParseError,
    UnsupportedFormatError,
    default_toolkit,
)

WORK = Path(tempfile.mkdtemp(prefix="doc_toolkit_test_"))
kit = default_toolkit
passed, failed = 0, []


def check(name: str, cond: bool, extra: str = "") -> None:
    global passed
    if cond:
        passed += 1
        print(f"  PASS {name}")
    else:
        failed.append(name)
        print(f"  FAIL {name} {extra}")


MD_SAMPLE = """# 项目说明

这是**第一段**介绍文字，包含行内标记。

## 功能列表

- 支持多格式提取
- 支持格式互转
  - 嵌套列表项

1. 第一步
2. 第二步

> 这是一段引用文字

```python
print("hello")
```

结尾段落。
"""

# ---- 1. 准备样例文件 -------------------------------------------------
md = WORK / "sample.md"
md.write_text(MD_SAMPLE, encoding="utf-8")

txt_gbk = WORK / "gbk_sample.txt"
txt_gbk.write_text("第一段：中文编码测试。\n\n第二段：GBK 编码段落。", encoding="gbk")

from docx import Document as NewDocx  # noqa: E402

docx_file = WORK / "sample.docx"
d = NewDocx()
d.add_heading("转换测试文档", level=1)
d.add_paragraph("这是正文第一段。")
d.add_heading("要点", level=2)
d.add_paragraph("第一个要点", style="List Bullet")
d.add_paragraph("第二个要点", style="List Bullet")
d.add_paragraph("收尾段落。")
d.save(str(docx_file))

from fpdf import FPDF  # noqa: E402

pdf_file = WORK / "sample.pdf"
p = FPDF()
p.add_page()
p.add_font("cjk", "", "C:/Windows/Fonts/msyh.ttc")
p.set_font("cjk", size=16)
p.multi_cell(0, 10, "PDF 转换测试", new_x="LMARGIN", new_y="NEXT")
p.set_font("cjk", size=11)
p.multi_cell(0, 6, "这是 PDF 正文段落，用于验证解析器。", new_x="LMARGIN", new_y="NEXT")
p.output(str(pdf_file))

print("== 1. 内容提取 ==")
for f in (md, txt_gbk, docx_file, pdf_file):
    try:
        text = kit.extract_text(f)
        check(f"extract_text {f.name}", len(text) > 5)
    except Exception as exc:
        check(f"extract_text {f.name}", False, repr(exc))

print("== 2. 结构保留（MD→IR）==")
doc = kit.parse(md)
types = [b.type.value for b in doc.blocks]
check("heading level1", any(t == "heading" for t in types))
check("unordered list", any(t == "list_item" for t in types))
check("code block", "code" in types)
check("quote", "quote" in types)

print("== 3. 格式互转 ==")
# 输出路径显式隔离，避免 with_suffix 与源文件同名互相覆盖。
cases = [
    (md, "docx", "out1.docx"), (md, "pdf", "out1.pdf"), (md, "txt", "out1.txt"),
    (docx_file, "markdown", "out2.md"), (docx_file, "pdf", "out2.pdf"), (docx_file, "txt", "out2.txt"),
    (pdf_file, "markdown", "out3.md"), (pdf_file, "docx", "out3.docx"), (txt_gbk, "docx", "out4.docx"),
]
for src, fmt, out_name in cases:
    try:
        out = kit.convert(src, fmt, WORK / out_name)
        check(f"convert {src.name} → {fmt}", out.is_file() and out.stat().st_size > 50)
    except Exception as exc:
        check(f"convert {src.name} → {fmt}", False, repr(exc))

print("== 4. 转换质量：DOCX→MD 结构回读 ==")
out_md = WORK / "out2.md"
back = kit.parse(out_md)
headings = [b for b in back.blocks if b.type.value == "heading"]
lists = [b for b in back.blocks if b.type.value == "list_item"]
check("标题保留", len(headings) >= 2 and headings[0].level == 1)
check("列表保留", len(lists) >= 2)

print("== 5. 错误提示 ==")
corrupt = WORK / "corrupt.docx"
corrupt.write_bytes(b"PK\x03\x04 garbage not a real docx" * 10)
try:
    kit.parse(corrupt)
    check("损坏 DOCX 报错", False, "未抛出异常")
except DocumentParseError as exc:
    check("损坏 DOCX 报错", "损坏" in str(exc) or "有效" in str(exc), str(exc))

fake_pdf = WORK / "corrupt.pdf"
fake_pdf.write_bytes(b"%PDF-1.4 broken content \xff\xfe" * 8)
try:
    kit.parse(fake_pdf)
    check("损坏 PDF 报错", False, "未抛出异常")
except DocumentParseError as exc:
    check("损坏 PDF 报错", bool(str(exc)), str(exc))

unknown = WORK / "unknown.xyz"
unknown.write_text("随便一点内容", encoding="utf-8")
try:
    kit.parse(unknown)
    check("未知扩展名报错", False)
except UnsupportedFormatError as exc:
    check("未知扩展名报错", "不支持" in str(exc), str(exc))

try:
    kit.extract_text(WORK / "missing.docx")
    check("文件不存在报错", False)
except Exception as exc:
    check("文件不存在报错", "不存在" in str(exc), str(exc))

print(f"\n结果：{passed} 通过，{len(failed)} 失败 {failed if failed else ''}")
print(f"样例目录：{WORK}")
sys.exit(1 if failed else 0)
